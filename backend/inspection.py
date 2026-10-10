import os
import fitz
import numpy as np   
import cv2 
import json 
import string 
import datetime
import time 
import io
import base64
from ultralytics import YOLO
from openpyxl import load_workbook
from utils.pipeline_utils import *
from utils.sheet import * 
from utils.insp_logger import user_logger, dev_logger, user_log_name
from utils.config import gen_cfg, dev_cfg
from utils.paddle_ocr import PaddleOCRCompat as vrdOCR
from transformers import ViTForImageClassification
from transformers import ViTImageProcessor

model_save_path = "models/prototypical_network.pth" 
prototypes_path = "models/prototypes.pth"  
classifier_files = [model_save_path, prototypes_path]
missing_classifier_files = [path for path in classifier_files if not os.path.isfile(path)]
if dev_cfg['enable_symbols'] and not missing_classifier_files:
    prototypical_model = load_model(PrototypicalNetwork(), model_save_path)
    prototypes = load_prototypes(prototypes_path)
else:
    prototypical_model = None
    prototypes = {}
    if missing_classifier_files:
        dev_cfg['enable_symbols'] = False
        warning = (
            "Care-symbol inspection disabled; missing classifier files: "
            + ", ".join(missing_classifier_files)
        )
        user_logger.warning(warning)
        dev_logger.warning(warning)

user_logger.info("INFO: Application starting, please wait...")
dev_logger.info(f"INFO: General Settings: {gen_cfg}")
dev_logger.info(f"INFO: Dev Settings: {dev_cfg}")
max_pages = gen_cfg['max_pages']
data_dir  = gen_cfg['data_dir']

dev_logger.info(f"INFO: Loading models...")

care_workbook = load_workbook(gen_cfg['care_workbook_path'])

fiber_workbook_ppta = load_workbook(gen_cfg['fiber_workbook_ppta_path'])
fiber_workbook_woppta = load_workbook(gen_cfg['fiber_workbook_woppta_path'])

part_workbook = load_workbook(gen_cfg['part_workbook_path'])
part_ext_workbook = load_workbook(gen_cfg['part_workbook_path_ext'])
symbols_workbook = load_workbook(gen_cfg['symbols_workbook_path'])
part_ext_workbook = part_ext_workbook.active 
part_workbook = part_workbook.active 
care_sheet = care_workbook.active

fiber_sheet_ppta = fiber_workbook_ppta.active
fiber_sheet_woppta = fiber_workbook_woppta.active
symbols_sheet = symbols_workbook.active
parts_sheet = {'Name': 'Part Position'  ,'main_sheet_range': part_workbook[gen_cfg['part_range'][0]:gen_cfg['part_range'][1]],  'extra_sheet_range': None}
parts_df    = preprocess_gt(table_2_df(sheet=part_workbook, table=parts_sheet))
parts_ext_sheet = {'Name': 'Part Position Extensive'  ,'main_sheet_range': part_ext_workbook[gen_cfg['part_ext_range'][0]:gen_cfg['part_ext_range'][1]],  'extra_sheet_range': None}
parts_ext_df = preprocess_gt(table_2_df(sheet=part_ext_workbook, table=parts_ext_sheet))
fiber_range_ppta = gen_cfg['fiber_range_ppta']
fiber_range_woppta = gen_cfg['fiber_range_woppta']
all_fiber_content_table_ppta = {'Name': 'all_fiber_content', 'main_sheet_range': fiber_sheet_ppta[fiber_range_ppta[0]: fiber_range_ppta[1]], 'extra_sheet_range': None}
all_fiber_content_df_ppta = preprocess_gt(table_2_df(sheet=fiber_sheet_ppta, table=all_fiber_content_table_ppta))
all_fiber_content_table_woppta = {'Name': 'all_fiber_content', 'main_sheet_range': fiber_sheet_woppta[fiber_range_woppta[0]: fiber_range_woppta[1]], 'extra_sheet_range': None}
all_fiber_content_df_woppta = preprocess_gt(table_2_df(sheet=fiber_sheet_woppta, table=all_fiber_content_table_woppta))
all_extra_care_instr_table = {'Name': 'all_extra_care_instructions', 'main_sheet_range': care_sheet[gen_cfg['care_range'][0]: gen_cfg['care_range'][1]], 'extra_sheet_range': None}
all_extra_care_instr_df = table_2_df(sheet=care_sheet, table=all_extra_care_instr_table)
all_care_instr_table = {'Name': 'all_care_instructions', 'main_sheet_range': care_sheet[gen_cfg['care_range'][0]: gen_cfg['care_range'][1]], 'extra_sheet_range': None}
all_care_instr_df = table_2_df(sheet=care_sheet, table=all_care_instr_table)
all_symbols = {'Name': 'all_symbols'  ,'main_sheet_range': symbols_sheet[gen_cfg['symbol_range'][0]:gen_cfg['symbol_range'][1]],  'extra_sheet_range': None}
all_symbols_df = preprocess_gt(table_2_df(sheet=symbols_sheet, table=all_symbols))

# Load US base size sheet
us_base_sheet = load_workbook(gen_cfg['us_size_path'], read_only=True, data_only=True)
us_base_sheet = us_base_sheet.active
us_base_data = us_base_sheet.values
us_base_columns = next(us_base_data)
us_size_df = pd.DataFrame(us_base_data, columns=us_base_columns)

# Load ADIDAS base size sheet 
adidas_base_sheet = load_workbook(gen_cfg['adidas_size_path'], read_only=True, data_only=True)
adidas_base_sheet = adidas_base_sheet.active
adidas_base_data = adidas_base_sheet.values
adidas_base_columns = next(adidas_base_data)
adidas_size_df = pd.DataFrame(adidas_base_data, columns=adidas_base_columns)

# Load English OCR
ocr_settings = json.load(open('configs/ocr_settings.json', '+r'))
vrd_en = vrdOCR(show_log=False)
vrd_search = vrdOCR(**ocr_settings)
vrd_unbatched = vrdOCR(se_angle_cls=True, lang='en',rec_batch_num=1, show_log=False)
vrd_latin = vrdOCR()

def run_pipeline(input_data, output_queue):
    ts_pipeline = time.time()
    try:
        global max_pages, data_dir, care_workbook, fiber_workbook_ppta, fiber_workbook_woppta
        global part_workbook, symbols_workbook, ocr_settings, vrd_en, vrd_latin
        global vit_symbol_processor, vit_symbol_classifier, scissors_detector, page_detector

        today = datetime.date.today()
        c = datetime.datetime.now()
        current_time = c.strftime('%H-%M-%S')
        item_name = input_data['item_name']
        
        input_folder =  f'inputs/{item_name}_{today}_{current_time}/'
        # os.makedirs(input_folder, exist_ok=True)

        excel_path = os.path.join(input_folder, f'{item_name}_{current_time}.xlsx')
        pdf_path   = os.path.join(input_folder, f'{item_name}_{current_time}.pdf')

        excel_status = decode_and_save(input_data['excel']['order_form'], path=excel_path)
        if excel_status != 0:
            return excel_status 

        pdf_status   = decode_and_save(input_data['pdf']['pdf_file'], path=pdf_path)
        if pdf_status != 0:
            return pdf_status 

        # Decode base64 and use BytesIO
        try:
            excel_base64 = input_data['excel']['order_form']
            excel_bytes = base64.b64decode(excel_base64)
            excel_io = io.BytesIO(excel_bytes)
            dev_logger.info("Excel file loaded into memory.")
        except Exception as e:
            dev_logger.info(f"Error decoding base64 content: {e}")
            raise

        try:
            metadata = input_data['ref']
            ref_status = 0
        except Exception as e: 
            ret_status = 1 
            defect_str = f"ERROR: Could not read REFERENCE file, please ensure correct inputs and try again."
            return send_pipeline_fail(defect_str)


        if excel_status != 1 and pdf_status != 1 and ref_status != 1:
            user_logger.info(f"INFO: Starting inspection for {item_name}...")
            dev_logger.info( f"INFO: Starting inspection for {item_name}...")

            items = [f'{item_name}']
            item_pngs = {item:{"img_list":[], "metadata":None} for item in items}
            for item in items:
                item_pngs[item]['metadata'] = metadata
                production_mode = item_pngs[item]['metadata']['production_mode']

                search_mode = item_pngs[item]['metadata'].get('searchmode',False)
                exsheet_names = item_pngs[item]['metadata'].get('sheet_name',[])
                #exsheet_names = ['IHT manual order form','UPC & CPSIA UPC label']
                column_names = item_pngs[item]['metadata'].get('column_names',[]) 
                singlekeyword = item_pngs[item]['metadata'].get('singlekeyword',[])


                if not search_mode:
                    vit_symbol_processor  = ViTImageProcessor.from_pretrained(gen_cfg['vit_model'])
                    vit_symbol_classifier = ViTForImageClassification.from_pretrained(gen_cfg['vit_model'])
                    vit_symbol_classifier.to('cuda')

                    scissors_detector = YOLO(gen_cfg['scissors_model'])
                    # if not production_mode:
                    page_detector = YOLO(gen_cfg['page_model'])

                ppta_status = item_pngs[item]['metadata']['ppta']

                item_pdf = fitz.open(pdf_path)
                for page_num in range(max_pages):
                    try:
                        page = item_pdf.load_page(page_num)
                        pixmap = page.get_pixmap(dpi=dev_cfg['pdf_dpi'])
                        img = pixmap.tobytes()
                        img = np.frombuffer(img, np.uint8)

                        img = cv2.imdecode(img,cv2.IMREAD_COLOR)
                        if production_mode:
                            img = cv2.rotate(img, cv2.ROTATE_90_CLOCKWISE)
                        item_pngs[item]["img_list"].append(img)
                    except:
                        break 
                
            results_folder = input_folder.replace('inputs', 'results')
            # os.makedirs(results_folder, exist_ok=True)
            input_folder = f'inputs/{item_name}_{today}_{current_time}/'
            # os.makedirs(input_folder, exist_ok=True)

            excel_path = os.path.join(input_folder, f'{item_name}_{current_time}.xlsx')
            if not search_mode:
                try:
                    order_form           = load_workbook(excel_path)
                except:
                    order_form = load_workbook(excel_path, data_only=True)

                order_form_READ_ONLY = load_workbook(excel_path, data_only=True)

                order_form_READ_ONLY = order_form_READ_ONLY.active
                order_form = order_form.active
            else:
                try:
                    order_form = pd.read_excel(excel_path, sheet_name=None) 
                except Exception as e:
                    dev_logger.info(f"Error reading Excel file with pandas: {e}")
                    raise

                sheet_names = list(order_form.keys())  # List all sheet names
                first_sheet_name = sheet_names[0]  # Use the first sheet
                order_form_READ_ONLY = order_form[first_sheet_name]

            if os.path.exists(pdf_path):
                os.remove(pdf_path)
                dev_logger.info('INFO: Artwork file loaded and deleted from disk')

            if os.path.exists(excel_path):
                os.remove(excel_path)
                dev_logger.info('INFO: Order form file loaded and deleted from disk')


            item_metadata = item_pngs[item]['metadata']
            fc_range      = item_metadata['fc_range']
            care_range    = item_metadata['care_range']
            excare_range  = item_metadata['excare_range']
            symbol_range  = item_metadata['symbol_range']

            size_range = ['B35', 'I64']
            size_langs = ['UK', 'US', 'F', 'D', 'I', 'E', 'CDN', 'J']
            size_translation='Adidas Size Translation'

            if len(fc_range) == 1 and '/' in fc_range[0]:
                fc_range = fc_range[0].split('/')

            if len(care_range) == 1 and '/' in care_range[0]:
                care_range = care_range[0].split('/')

            if len(excare_range) == 1 and '/' in excare_range[0]:
                excare_range = excare_range[0].split('/')

            if len(symbol_range) == 1 and '/' in symbol_range[0]:
                symbol_range = symbol_range[0].split('/')

            fc_check = False
            care_check = False 
            excare_check = False 
            symbol_check = False 
            size_check = False 

            if ppta_status:
                all_fiber_content_df = all_fiber_content_df_ppta
            else: 
                all_fiber_content_df = all_fiber_content_df_woppta
                
            if len(fc_range) and dev_cfg['enable_fc'] and not search_mode:
                fc_check = True  
            else:
                fc_read_status = True 
                dev_logger.info(f"WARNING: Forcefully set fiber content read status to True. Dev Enabled: {dev_cfg['enable_fc']}, Empty Range: {fc_range}")

            if len(care_range) and dev_cfg['enable_care'] and not search_mode:
                care_check = True
            else:
                dev_logger.info(f"WARNING: Forcefully set care instr read status to True. Dev Enabled: {dev_cfg['enable_care']}, Empty Range: {care_range}")
                care_read_status = True 
            if len(excare_range) and dev_cfg['enable_excare'] and not search_mode:
                excare_check = True 
            else:
                excare_read_status = True
                dev_logger.info(f"WARNING: Forcefully set extra care instr read status to True. Dev Enabled: {dev_cfg['enable_excare']}, Empty Range: {excare_range}")
            if len(symbol_range) and dev_cfg['enable_symbols'] and not search_mode:
                symbol_check = True
            else: 
                care_symbol_read_status = True
                dev_logger.info(f"WARNING: Forcefully set care symbol read status to True. Dev Enabled: {dev_cfg['enable_symbols']}, Empty Range: {symbol_range}")
            
            if len(size_range) and dev_cfg['enable_sizes'] and not search_mode:
                size_check=True 
            else: 
                size_read_status = True 
                dev_logger.info(f"WARNING: Forcefully set size read status to True. Dev Enabled: {dev_cfg['enable_sizes']}, Empty Range: {size_range}")

            # Load fiber content
            if fc_check:
                fiber_content_table     = {'Name': 'fiber_content' ,'main_sheet_range': order_form[fc_range[0]:fc_range[1]], 'extra_sheet_range': None}
                fc_handle_out = fc_table_handler(sheet=order_form, table=fiber_content_table)
                try: 
                    fc_read_status, fc_gt, fiber_content_df = fc_handle_out[0], fc_handle_out[1], fc_handle_out[2]
                except:
                    return fc_handle_out
                
            if excare_check:
                # Load extra care instructions 
                extra_care_instr_table     = {'Name': 'extra_care_instructions' ,'main_sheet_range': order_form[excare_range[0]:excare_range[1]], 'extra_sheet_range': None}
                excare_handle_out =  excare_instr_table_handler(sheet=order_form, table=extra_care_instr_table, ref_sheet=all_extra_care_instr_df)
                try: 
                    excare_read_status, excare_instr_dict = excare_handle_out[0], excare_handle_out[1]
                except:
                    return excare_handle_out 
                
            if care_check:
                # Load care instructions & symbols 
                care_instr_table     = {'Name': 'care_instructions' ,'main_sheet_range': order_form_READ_ONLY[care_range[0]:care_range[1]], 'extra_sheet_range': None}
                care_handle_out = care_instr_table_handler(sheet=order_form, table=care_instr_table, ref_sheet=all_care_instr_df)        
                try: 
                    care_read_status, care_instr_dict =  care_handle_out[0], care_handle_out[1]
                except:
                    return care_handle_out 
                
            if symbol_check:
                try:
                    care_symbols_table   = {'Name': 'care_symbols' ,'main_sheet_range': order_form[symbol_range[0]:symbol_range[1]], 'extra_sheet_range': None}
                    care_symbols_df = table_2_df_order(sheet=order_form, table=care_symbols_table)
                    care_symbols    = preprocess_gt(care_symbols_df).iloc[0]
                    if len(care_symbols):
                        care_symbol_read_status = True
                except: 
                    defect_str = f'ERROR: Could not correctly read care symbols from order form, please check your inputs / select valid table ranges.'
                    return send_pipeline_fail(defect_str)
            
            if size_check: 
                size_table     = {'Name': 'size_table' ,'main_sheet_range': order_form[size_range[0]:size_range[1]], 'extra_sheet_range': None}
                if size_translation == 'US Size Translation':
                    size_read_status, size_gt = us_size_handler(order_form=order_form,
                                                base_df=us_size_df,
                                                size_table=size_table,
                                                size_langs=size_langs)
                    
                elif size_translation == 'Adidas Size Translation':
                    size_read_status, size_gt = adidas_size_handler(order_form=order_form,
                                                base_df=adidas_size_df,
                                                size_table=size_table,
                                                size_langs=size_langs)
                
                if not size_read_status:
                    defect_str = f'ERROR: Something went wrong while reading the size table. Please make sure your configuration/base file is correct.'
                    dev_logger.info(f'{defect_str},\n{size_gt}\n{traceback.format_exc()}')
                    user_logger.info(f'{defect_str}')
                    return send_pipeline_fail(defect_str)
                
            if fc_read_status and excare_read_status and care_read_status and care_symbol_read_status:
                # Settings 
                if item_metadata['order_market'][0] == 'EU':
                    order_market = ' (EU)'
                elif item_metadata['order_market'][0] == 'US':
                    order_market = ' (US)'
                else: 
                    order_market = dev_cfg['default_order_market']#' (EU)'
                    user_logger.warning(f"WARNING: Could not get Order Market from reference JSON. Proceeding with default: {order_market}")
                    dev_logger.warning( f"WARNING: Could not get Order Market from reference JSON. Proceeding with default: {order_market}")

                if len(item_pngs[item]['metadata']['care_langs']) == 1 and '/' in item_pngs[item]['metadata']['care_langs'][0]: 
                    item_pngs[item]['metadata']['care_langs'] = item_pngs[item]['metadata']['care_langs'][0].split('/')

                if len(item_pngs[item]['metadata']['fc_langs']) == 1 and '/' in item_pngs[item]['metadata']['fc_langs'][0]: 
                    item_pngs[item]['metadata']['fc_langs'] = item_pngs[item]['metadata']['fc_langs'][0].split('/')

                if len(item_pngs[item]['metadata']['excare_langs']) == 1 and '/' in item_pngs[item]['metadata']['excare_langs'][0]: 
                    item_pngs[item]['metadata']['excare_langs'] = item_pngs[item]['metadata']['excare_langs'][0].split('/')

                blacklisted_langs = {}
                blacklisted_langs['missing'] = set()
                blacklisted_langs['extra'] = set()
                blacklisted_langs['unsupported'] = set()#{'CN', 'AE'} #| {'EST', 'SLO', 'RU', 'KZ', 'SRB', 'N', 'HR', 'BG', 'SF', 'DK', 'TR', 'NL', 'LV', 'RO', 'AE', 'S', 'HU', 'CZ', 'GR', 'PL', 'LT', 'I', 'SK'}
                all_langs = set(get_possible_langs(item_pngs[item]['metadata']))
                latin_langs = gen_cfg['latin_langs']
                
                for gt_lang in all_langs: 
                    if gt_lang not in latin_langs:
                        user_logger.warning(f"WARNING: Blacklisted unsupported language from inspection: {gt_lang}")
                        dev_logger.warning(f"WARNING: Blacklisted unsupported language from inspection: {gt_lang}")
                        blacklisted_langs['unsupported'].add(gt_lang)

                care_langs_gt   = {lang for lang in set(item_pngs[item]['metadata']['care_langs'])   if lang not in blacklisted_langs['unsupported'] }
                fc_langs_gt     = {lang for lang in set(item_pngs[item]['metadata']['fc_langs'])     if lang not in blacklisted_langs['unsupported'] }
                excare_langs_gt = {lang for lang in set(item_pngs[item]['metadata']['excare_langs']) if lang not in blacklisted_langs['unsupported'] }

                reverse_map_fc   = gen_cfg['reverse_map_fc']
                reverse_map_part = gen_cfg['reverse_map_part']

                care_symbol_reverse_map = {}
                for symbol_idx in range(len(all_symbols_df)):
                    symbol_name = all_symbols_df.iloc[symbol_idx].tolist()[0]
                    care_symbol_reverse_map[symbol_name] = symbol_idx

                possible_care_symbols = gen_cfg['possible_care_symbols']
                fiber_content_gt = {}
                if fc_check:
                    
                    langwise_contents = {}
                    container = {}
                    for fc_lang in fc_langs_gt: ### 'CN' ignored cus not in excel
                        if fc_lang in latin_langs:
                            use_decode=True 

                        fc_lang_org = fc_lang
                        if fc_lang == 'UK' or fc_lang == 'US':
                            fc_lang = 'GB' 
                        if fc_lang == 'HR':
                            fc_lang = 'CROATIA (HR)'
                        if fc_lang == 'KZ':
                            fc_lang = 'KAZAKHSTAN'
                        if fc_lang == 'SRB':
                            fc_lang = 'SRBIJA'
                        if fc_lang == 'IND':
                            try:
                                possible_contents = all_fiber_content_df[fc_lang]
                            except:
                                try: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' in FC spreadsheet, trying 'INDONESIAN'")
                                    fc_lang = 'INDONESIAN'
                                    possible_contents = all_fiber_content_df[fc_lang]
                                    dev_logger.info("INFO:'INDONESIAN' was found in FC sheet.")
                                except: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in FC spreadsheet, lang blacklisted.")
                                    blacklisted_langs['unsupported'].add(fc_lang)
                                    user_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in FC spreadsheet, lang blacklisted.")

                        if fc_lang == 'GB' or fc_lang == 'E' or fc_lang == 'P' or fc_lang == 'F':
                            fc_lang += order_market

                        langwise_contents = {}
                        possible_contents = all_fiber_content_df[fc_lang]
                        eu_english_contents = all_fiber_content_df['GB (EU)']
                        us_english_contents = all_fiber_content_df['GB (US)']

                        possible_content_types = parts_df[fc_lang]

                        eu_english_parts = parts_ext_df['GB (EU)']
                        us_english_parts = parts_ext_df['GB (US)']

                        for content_type, content_dict in fc_gt.items():
                            try:
                                content_type_idx, content_type_dist = find_fc_idx(us_english_parts, eu_english_parts, content_type, max_distance=dev_cfg['default_str_thresh'])
                            except:
                                dev_logger.info(f'WARNING: Could not find part position {content_type} in extended part position sheet in US+EU. Retrying with US+US')
                            try:
                                content_type_idx, content_type_dist = find_fc_idx(us_english_parts, us_english_parts, content_type, max_distance=dev_cfg['default_str_thresh'])
                            except:
                                dev_logger.info(f'WARNING: Could not find part position {content_type} in extended part position sheet in US+US. Retrying with EU+EU')
                            try:
                                content_type_idx, content_type_dist = find_fc_idx(eu_english_parts, eu_english_parts, content_type, max_distance=dev_cfg['default_str_thresh'])
                            except:
                                defect_str = f'ERROR: Could not find part position {content_type} in extended part position sheet. Please ensure you have selected correct inputs.'
                                user_logger.info(defect_str)
                                dev_logger.info(defect_str)
                                return send_pipeline_fail(defect_str)
                            
                            try: 
                                content_type = parts_ext_df[fc_lang][content_type_idx]
                            except: 
                                defect_str = f'ERROR: Could not find part position {content_type} in extended part position sheet for language: {fc_lang}. Please ensure you have selected correct inputs.'
                                user_logger.info(defect_str)
                                dev_logger.info(defect_str)
                                return send_pipeline_fail(defect_str)
                            
                            container = {}
                            if use_decode:
                                content_type = unidecode(content_type)

                            for material, material_percentage in content_dict.items():
                                split_material=False
                                if '///' in material:
                                    
                                    split_material=True
                                    material_unjoined, actual_material = material.split('///')[0], material.split('///')[-1]
                                    for joinable_type in gen_cfg['joinable_material_types']:
                                        if joinable_type in material:
                                            material_unjoined = material_unjoined.replace(joinable_type, '')
                                            break 

                                    material_sheet_idx, material_dist = find_fc_idx(eu_english_contents, us_english_contents, material_unjoined, max_distance=dev_cfg['default_str_thresh'])
                                    joined_material_sheet_idx, joined_material_dist = find_fc_idx(eu_english_contents, us_english_contents, material_unjoined+joinable_type, max_distance=dev_cfg['default_str_thresh'])
                                    try: 
                                    
                                        unjoined_container_key = possible_contents[material_sheet_idx]
                                        joined_container_key = possible_contents[joined_material_sheet_idx]
                                        if '(' in joined_container_key:
                                            refined_container_key = joined_container_key.replace('(', f'({extract_numeric_value(actual_material)}% ')
                                        else:
                                            diff_index = None
                                            min_len = min(len(unjoined_container_key), len(joined_container_key))
                                            for i in range(min_len):
                                                if unjoined_container_key[i] != joined_container_key[i]:
                                                    diff_index = i
                                                    break
                                            if diff_index is None:
                                                diff_index = min_len
                                            refined_container_key = unjoined_container_key[:diff_index] + f'({extract_numeric_value(actual_material)}%' + joined_container_key[diff_index:] + ')'
                                        if use_decode:
                                            container_key = unidecode(refined_container_key)
                                        else:
                                            container_key = refined_container_key
                                    except Exception as e:
                                            defect_str = f"ERROR: Could not find {material} equivalent in lang {fc_lang} in supplementary fiber content spreadsheet. Please ensure you have selected the correct file."
                                            return send_pipeline_fail(defect_str)
                                    
                                if use_decode:
                                    if not split_material:
                                        material_sheet_idx, material_dist = find_fc_idx(eu_english_contents, us_english_contents, material, max_distance=dev_cfg['default_str_thresh'])


                                        if material_sheet_idx is None: 

                                            
                                            defect_str = f"ERROR: Could not find {material} in supplementary fiber content spreadsheet. Please ensure you have selected the correct file."
                                            return send_pipeline_fail(defect_str)
                                        else: 
                                            try: 
                                                container_key = possible_contents[material_sheet_idx]
                                            except Exception as e:
                                                defect_str = f"ERROR: Could not find {material} equivalent in lang {fc_lang} in supplementary fiber content spreadsheet. Please ensure you have selected the correct file."
                                                return send_pipeline_fail(defect_str)
                                    if container_key: 
                                        
                                        container[unidecode(container_key)] = material_percentage
                                    else: 
                                        defect_str = f"ERROR: Could not find {material} in supplementary fiber content spreadsheet. Please ensure you have selected the correct file."
                                        return send_pipeline_fail(defect_str)
                                    try: 
                                        langwise_contents[content_type] = container
                                    except Exception as e:
                                        defect_str = f"ERROR: Could not find {content_type} in supplementary part position spreadsheet. Please ensure you have selected the correct file."
                                        return send_pipeline_fail(defect_str)

                                else: 
                                    container[possible_contents[reverse_map_fc[material]]] = material_percentage
                                    langwise_contents[possible_content_types[reverse_map_part[content_type]]] = container
                                fiber_content_gt[fc_lang_org] = langwise_contents
                extra_care_gt = {}
                
                if excare_check:
                    
                    for care_lang in excare_langs_gt: 
                        if care_lang in latin_langs:
                            use_decode = True 

                        care_lang_org = care_lang
                        if care_lang == 'UK' or care_lang == 'US':
                            care_lang = 'GB' 
                        if care_lang == 'HR':
                            care_lang = 'CROATIA (HR)'
                        if care_lang == 'KZ':
                            care_lang = 'KAZAKHSTAN'
                        if care_lang == 'SRB':
                            care_lang = 'SERBIJA'
                        if care_lang == 'IND':
                            try:
                                instruction_all_lang = list(excare_instr_dict.values())[0]
                                instruction_langwise = instruction_all_lang[care_lang]
                            except:
                                try: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' in care spreadsheet, trying 'INDONESIAN'")
                                    care_lang = 'INDONESIAN'
                                    instruction_langwise = instruction_all_lang[care_lang]
                                    dev_logger.info("INFO:'INDONESIAN' was found in care sheet.")
                                except: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in extra care spreadsheet, lang blacklisted.")
                                    blacklisted_langs['unsupported'].add(care_lang)
                                    user_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in extra care spreadsheet, lang blacklisted.")

                        if care_lang == 'GB' and order_market == ' (EU)':
                            care_lang = care_lang 
                        if care_lang == 'GB' and order_market == ' (US)':
                            care_lang += order_market
                        if care_lang == 'E' or care_lang == 'P' or care_lang == 'F':
                            care_lang += order_market
                        if care_lang == 'P (US)':
                            care_lang = 'P (US) - (Portuguese Brazil) '  

                        extra_care_gt[care_lang_org] = [] 
                        for instr_id, instruction_all_lang in excare_instr_dict.items():
                            if instr_id != 462:
                                instruction_langwise = instruction_all_lang[care_lang].item() 
                            else: 
                                instruction_langwise = 'EXCLUSIVE OF TRIMMING AND DECORATION'
                            if instruction_langwise is not None: 
                                if use_decode:
                                    extra_care_gt[care_lang_org].append(unidecode(instruction_langwise))
                                else:
                                    extra_care_gt[care_lang_org].append(instruction_langwise)

                care_instr_gt = {}
                if care_check:
                    
                    for care_lang in care_langs_gt: 
                        if care_lang in latin_langs:
                            use_decode = True 
                        care_lang_org = care_lang
                        if care_lang == 'UK' or care_lang == 'US':
                            care_lang = 'GB' 
                        if care_lang == 'HR':
                            care_lang = 'CROATIA (HR)'
                        if care_lang == 'KZ':
                            care_lang = 'KAZAKHSTAN'
                        if care_lang == 'SRB':
                            care_lang = 'SERBIJA'
                        if care_lang == 'IND':
                            try:
                                instruction_langwise = instruction_all_lang[care_lang]
                            except:
                                try: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' in care spreadsheet, trying 'INDONESIAN'")
                                    care_lang = 'INDONESIAN'
                                    instruction_langwise = instruction_all_lang[care_lang]
                                    dev_logger.info("INFO:'INDONESIAN' was found in care sheet.")
                                except: 
                                    dev_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in care spreadsheet, lang blacklisted.")
                                    blacklisted_langs['unsupported'].add(care_lang)
                                    user_logger.info("WARNING: Could not find lang 'IND' or 'INDONESIAN' in care spreadsheet, lang blacklisted.")
                        if care_lang == 'GB' and order_market == ' (EU)':
                            care_lang = care_lang 
                        if care_lang == 'GB' and order_market == ' (US)':
                            care_lang += order_market
                        if care_lang == 'E' or care_lang == 'P' or care_lang == 'F':
                            care_lang += order_market
                        if care_lang == 'P (US)':
                            care_lang = 'P (US) - (Portuguese Brazil) '  
                        care_instr_gt[care_lang_org] = [] 
                        for instr_id, instruction_all_lang in care_instr_dict.items(): 

                            if instr_id != 462:
                                instruction_langwise = instruction_all_lang[care_lang].item() 
                            else: 
                                instruction_langwise = 'EXCLUSIVE OF TRIMMING AND DECORATION'
                            if instruction_langwise is not None: 
                                if use_decode:
                                    care_instr_gt[care_lang_org].append(unidecode(instruction_langwise))
                                else:
                                    care_instr_gt[care_lang_org].append(instruction_langwise)
                care_symbols_gt = {}
                if symbol_check:
                    for possible_care_symbol in possible_care_symbols:
                        try:
                            hold = care_symbols[possible_care_symbol]
                            if hold is not None:
                                symbol_details = all_symbols_df.iloc[care_symbol_reverse_map[hold]]
                                care_instruction = symbol_details['UK']
                                care_symbols_gt[possible_care_symbol] = {'symbol_id':hold, 'symbol_instr': care_instruction}
                        except:
                            continue
                gt = {
                        'fc_gt': fiber_content_gt, 
                        'care_gt': care_instr_gt, 
                        'excare_gt': extra_care_gt, 
                        'symbols_gt': care_symbols_gt,
                        'pdf_image': item_pngs[item]['img_list']
                    }
                page_bboxes = {}
                
                dev_logger.info(f"INFO: FC GT: {gt['fc_gt']} Care GT: {gt['care_gt']} ExCare GT: {gt['excare_gt']} Symbols GT: {gt['symbols_gt']}")
                torch.cuda.empty_cache()
                size_gt = {}
                if production_mode:
                    inspection_dict = production_mode_function(
                    item_pngs=item_pngs,
                    results_folder=results_folder,
                    dev_logger=dev_logger,
                    user_logger=user_logger,
                    vrd_en=vrd_en,
                    encode_img=encode_img,
                    encode_str=encode_str,
                    user_log_name=user_log_name,
                    ts_pipeline=ts_pipeline,
                    dev_cfg=dev_cfg,
                    fc_check=fc_check,
                    care_check=care_check,
                    excare_check=excare_check,
                    symbol_check=symbol_check,
                    care_langs_gt=care_langs_gt,
                    excare_langs_gt=excare_langs_gt,
                    all_langs=all_langs,
                    blacklisted_langs=blacklisted_langs,
                    fiber_content_gt=fiber_content_gt,
                    care_instr_gt=care_instr_gt,
                    extra_care_gt=extra_care_gt,
                    page_detector=page_detector,
                    scissors_detector=scissors_detector,
                    gt=gt,
                    page_bboxes=page_bboxes,
                    vrd_latin=vrd_latin, 
                    fc_langs_gt=fc_langs_gt, 
                    prototypical_model=prototypical_model, 
                    prototypes=prototypes, 
                    care_symbols_gt=care_symbols_gt,
                    singlekeyword=singlekeyword,
                    size_gt=size_gt
                )

                    return inspection_dict 


                elif search_mode:
                    inspection_dict = search_mode_function(
                    excel_io=excel_io,
                    sheet_names=exsheet_names,
                    column_names=column_names,
                    item_pngs=item_pngs,
                    results_folder=results_folder,
                    dev_logger=dev_logger,
                    user_logger=user_logger,
                    vrd_unbatched=vrd_unbatched,
                    vrd_en=vrd_search,
                    encode_img=encode_img,
                    encode_str=encode_str,
                    user_log_name=user_log_name,
                    ts_pipeline=ts_pipeline,
                    gt=gt,
                    item=item,
                    singlekeyword=singlekeyword
                )
                    return inspection_dict

                                        
                
                else:
                    inspection_dict = layout_mode_function(
                    item_pngs=item_pngs,
                    results_folder=results_folder,
                    dev_logger=dev_logger,
                    user_logger=user_logger,
                    vrd_en=vrd_en,
                    encode_img=encode_img,
                    encode_str=encode_str,
                    user_log_name=user_log_name,
                    ts_pipeline=ts_pipeline,
                    dev_cfg=dev_cfg,
                    fc_check=fc_check,
                    care_check=care_check,
                    excare_check=excare_check,
                    symbol_check=symbol_check,
                    care_langs_gt=care_langs_gt,
                    excare_langs_gt=excare_langs_gt,
                    all_langs=all_langs,
                    blacklisted_langs=blacklisted_langs,
                    fiber_content_gt=fiber_content_gt,
                    care_instr_gt=care_instr_gt,
                    extra_care_gt=extra_care_gt,
                    page_detector=page_detector,
                    scissors_detector=scissors_detector,
                    gt=gt,
                    page_bboxes=page_bboxes,
                    vrd_latin=vrd_latin, 
                    fc_langs_gt=fc_langs_gt, 
                    prototypical_model=prototypical_model, 
                    prototypes=prototypes, 
                    care_symbols_gt=care_symbols_gt,
                    string=string,
                    item=item,
                    max_pages=max_pages,
                    singlekeyword=singlekeyword
                )
                    return inspection_dict
            
    except Exception as pipeline_fail:
        defect_str = f"ERROR 999: Something went wrong during inspection, please check your inputs/contact support."
        dev_logger.info(f"ERROR 999: Something went wrong during inspection, please check your inputs/contact support.\n{traceback.format_exc()}")
        return send_pipeline_fail(defect_str)
        
                    
